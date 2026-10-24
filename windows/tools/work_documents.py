"""Generate real Office/CSV files without controlling the desktop UI.

New files only. Paths are restricted by tools.paths.resolve_path.
"""
import csv
from pathlib import Path

from tools.paths import resolve_path
from tools.registry import register

MAX_ROWS = 3000
MAX_COLUMNS = 50
MAX_TEXT = 40000


def _destination(value, suffix):
    path = resolve_path(value)
    if path.suffix.lower() != suffix:
        raise ValueError("Expected a " + suffix + " destination.")
    if path.exists():
        raise FileExistsError("File already exists; choose another filename: " + str(path))
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def _columns_and_rows(columns, rows):
    if not isinstance(columns, list) or not 1 <= len(columns) <= MAX_COLUMNS:
        raise ValueError("columns must be a list with 1 to 50 entries.")
    headers = [str(x)[:150] for x in columns]
    if any(not x.strip() for x in headers):
        raise ValueError("Column names cannot be empty.")
    if not isinstance(rows, list) or len(rows) > MAX_ROWS:
        raise ValueError("rows must be a list of at most 3000 records.")
    records = []
    for row in rows:
        if not isinstance(row, list) or len(row) != len(headers):
            raise ValueError("Every row must be a list matching the columns.")
        records.append([_safe_value(value) for value in row])
    return headers, records


def _safe_value(value):
    if value is None or isinstance(value, (int, float, bool)):
        return value
    text = str(value)[:2000]
    # Prevent values from being interpreted as spreadsheet formulas.
    if text.startswith(("=", "+", "-", "@", "\t", "\r")):
        text = "'" + text
    return text


@register(
    "create_spreadsheet",
    "Create a genuine Excel .xlsx workbook with a formatted table. Never overwrites. "
    'arguments={path:"documents/report.xlsx",columns:["Item","Cost"],rows:[["Paper",5],["Ink",12]],sheet_name?}',
)
def create_spreadsheet(path, columns, rows, sheet_name="Sheet1"):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    headers, records = _columns_and_rows(columns, rows)
    target = _destination(path, ".xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = str(sheet_name or "Sheet1")[:31]
    ws.append(headers)
    for record in records:
        ws.append(record)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(records) + 1}"
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="1D3557")
        cell.alignment = Alignment(wrap_text=True)
    for index, name in enumerate(headers, 1):
        sample = [len(str(name))] + [len(str(row[index - 1])) for row in records[:150]]
        ws.column_dimensions[get_column_letter(index)].width = min(55, max(12, max(sample) + 2))
    wb.save(target)
    return f"Created Excel spreadsheet: {target} ({len(records)} rows)"


@register(
    "create_csv",
    'Create a CSV table for Excel or other apps; never overwrites. arguments={path:"documents/data.csv",columns:["Name","Value"],rows:[["A",1]]}',
)
def create_csv(path, columns, rows):
    headers, records = _columns_and_rows(columns, rows)
    target = _destination(path, ".csv")
    # utf-8-sig makes international text work better in Excel on Windows.
    with target.open("x", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        writer.writerows(records)
    return f"Created CSV file: {target} ({len(records)} rows)"


@register(
    "create_word_document",
    'Create a genuine Word .docx file with heading and paragraphs; never overwrites. '
    'arguments={path:"documents/proposal.docx",title:"Proposal",content:"Intro...",sections?:[{"heading":"Scope","body":"Details..."}]}',
)
def create_word_document(path, title, content="", sections=None):
    from docx import Document
    target = _destination(path, ".docx")
    if len(str(content)) > MAX_TEXT:
        raise ValueError("Document is too long.")
    doc = Document()
    doc.add_heading(str(title)[:250], 0)
    for paragraph in str(content).split("\n"):
        if paragraph.strip():
            doc.add_paragraph(paragraph)
    if sections is not None:
        if not isinstance(sections, list) or len(sections) > 40:
            raise ValueError("sections must be a list of up to 40 entries.")
        for section in sections:
            if not isinstance(section, dict):
                raise ValueError("Each section must contain heading and body.")
            doc.add_heading(str(section.get("heading", ""))[:180], level=1)
            text = str(section.get("body", ""))
            if len(text) > MAX_TEXT:
                raise ValueError("Section is too long.")
            for paragraph in text.split("\n"):
                if paragraph.strip():
                    doc.add_paragraph(paragraph)
    doc.save(target)
    return f"Created Word document: {target}"


@register(
    "create_presentation",
    'Create a genuine PowerPoint .pptx file; never overwrites. '
    'arguments={path:"documents/briefing.pptx",title:"Briefing",slides:[{"title":"Goals","bullets":["Grow","Improve"]}]}',
)
def create_presentation(path, title, slides):
    from pptx import Presentation
    target = _destination(path, ".pptx")
    if not isinstance(slides, list) or not 1 <= len(slides) <= 25:
        raise ValueError("slides must contain 1 to 25 entries.")
    prs = Presentation()
    first = prs.slides.add_slide(prs.slide_layouts[0])
    first.shapes.title.text = str(title)[:200]
    if len(first.placeholders) > 1:
        first.placeholders[1].text = "Created with J.A.R.V.I.S."
    for item in slides:
        if not isinstance(item, dict):
            raise ValueError("Each slide must contain title and bullets.")
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = str(item.get("title", ""))[:200]
        bullets = item.get("bullets", [])
        if not isinstance(bullets, list) or len(bullets) > 15:
            raise ValueError("Each slide supports up to 15 bullet points.")
        frame = slide.placeholders[1].text_frame
        frame.clear()
        for index, point in enumerate(bullets):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = str(point)[:400]
            paragraph.level = 0
    prs.save(target)
    return f"Created PowerPoint presentation: {target} ({len(slides) + 1} slides)"
