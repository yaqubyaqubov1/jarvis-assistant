"""Smoke tests for JARVIS work files and unsent email previews.

These tests never contact Gmail and never send email.
"""
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from tools.work_documents import (
    create_spreadsheet, create_csv, create_word_document, create_presentation,
)
from tools.gmail import _make_draft, draft_preview


class WorkFileTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory(dir=Path.home())
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)

    def test_xlsx_and_no_overwrite(self):
        from openpyxl import load_workbook
        path = self.folder / "book.xlsx"
        result = create_spreadsheet(str(path), ["Name", "Value"], [["A", 4], ["=SUM(1,2)", 9]])
        self.assertIn("Created Excel", result)
        wb = load_workbook(path)
        self.assertEqual(wb.active["A2"].value, "A")
        self.assertEqual(wb.active["A3"].data_type, "s")
        with self.assertRaises(FileExistsError):
            create_spreadsheet(str(path), ["X"], [[1]])

    def test_csv_escaping(self):
        path = self.folder / "table.csv"
        create_csv(str(path), ["Item"], [["=HYPERLINK(\"bad\")"]])
        text = path.read_text(encoding="utf-8-sig")
        self.assertIn("'=HYPERLINK", text)

    def test_docx(self):
        from docx import Document
        path = self.folder / "letter.docx"
        create_word_document(str(path), "Letter", "Hello world")
        self.assertEqual(Document(path).paragraphs[0].text, "Letter")

    def test_pptx(self):
        from pptx import Presentation
        path = self.folder / "brief.pptx"
        create_presentation(str(path), "Briefing", [{"title": "One", "bullets": ["A"]}])
        self.assertEqual(len(Presentation(path).slides), 2)


class EmailSafetyTests(unittest.TestCase):
    def test_email_draft_is_unsent_preview(self):
        result = _make_draft("test@example.org", "Hello", "Message body")
        self.assertIn("NOT sent", result)
        self.assertIn("Message body", draft_preview("last"))

    def test_reject_invalid_email(self):
        with self.assertRaises(ValueError):
            _make_draft("invalid-address", "Test", "Body")


if __name__ == "__main__":
    unittest.main()
